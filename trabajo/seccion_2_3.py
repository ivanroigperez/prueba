from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
d=Document("Trabajo_Taburete.docx")
YEL=WD_COLOR_INDEX.YELLOW

# ================= REVISIÓN DE LA ETAPA 1 =================
for p in d.paragraphs:
    if p.text.startswith("El problema es que estas alturas"):
        for r in p.runs: r.font.highlight_color=YEL      # dato de alcance P5 pendiente de confirmar
    if p.text.startswith("Para fijar la altura necesaria"):
        p.text=("Para fijar la altura necesaria se ha hecho un cálculo sencillo a partir de los datos del apartado anterior. "
        "Si el objetivo es llegar a 2,30 m y el alcance de una usuaria de percentil 5 es de unos 1,85 m, la plataforma superior "
        "tiene que estar como mínimo a 0,45 m del suelo. Por comodidad se toma una altura de 0,48 m, repartida en dos peldaños "
        "iguales de 0,24 m, ya que la norma UNE-EN 14183 exige que la distancia entre peldaños sea la misma que la del suelo al "
        "primer peldaño (Ilustración 3). Esta altura coincide prácticamente con la de los modelos de referencia del mercado, "
        "como el Hailo K20, con plataforma a 46 cm.")
    if p.text.startswith("Las propiedades traducen"):
        p.text=p.text+" Los valores de peso y de espesor plegado se han endurecido después del estudio de mercado (apartado 2.2), porque los valores iniciales quedaban por debajo de lo que ya ofrece la competencia."
tp=d.tables[1]
for row in tp.rows:
    c=row.cells
    if c[0].text=="P2": c[2].paragraphs[0].runs[0].text="Superar los ensayos de resistencia y de deslizamiento lateral de la UNE-EN 14183; base de apoyo más ancha que la plataforma"
    if c[0].text=="P5": c[2].paragraphs[0].runs[0].text="Peso total ≤ 3,5 kg"
    if c[0].text=="P6": c[2].paragraphs[0].runs[0].text="Espesor plegado ≤ 50 mm"
    if c[0].text=="P12": c[2].paragraphs[0].runs[0].text="Precio de venta objetivo ≤ 45 € IVA incluido (ver apartado 2.1)"

# quitar bibliografía para reponerla al final
idx=[i for i,p in enumerate(d.paragraphs) if p.text=="Bibliografía"][0]
old=d.paragraphs[idx:]; bib=[p.text for p in old[1:]]
for p in old: p._p.getparent().remove(p._p)

# ================= utilidades =================
def P(*segs,bold=False,style=None):
    p=d.add_paragraph(style=style) if style else d.add_paragraph()
    for s in segs:
        if isinstance(s,tuple): t,h=s
        else: t,h=s,False
        r=p.add_run(t); r.bold=bold
        if h: r.font.highlight_color=YEL
    return p
def B(t,h=False): return P((t,h),style="List Bullet")
def N(t,h=False): return P((t,h),style="List Number")
def H(t,l): d.add_heading(t,l)
def cap(t):
    c=d.add_paragraph(); c.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=c.add_run(t); r.italic=True; r.font.size=Pt(9)
def fig(path,t,w=15):
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(path,width=Cm(w)); cap(t)
def shade(cell,color):
    tcPr=cell._tc.get_or_add_tcPr(); s=OxmlElement("w:shd"); s.set(qn("w:val"),"clear"); s.set(qn("w:color"),"auto"); s.set(qn("w:fill"),color); tcPr.append(s)
def tabla(cab,filas,anchos,t,fs=8.5):
    tb=d.add_table(rows=1,cols=len(cab)); tb.style="Table Grid"
    for i,c in enumerate(cab):
        cell=tb.rows[0].cells[i]; cell.text=""; r=cell.paragraphs[0].add_run(c); r.bold=True; r.font.size=Pt(fs+0.5); r.font.color.rgb=RGBColor(255,255,255); shade(cell,"C0392B")
    for f in filas:
        cells=tb.add_row().cells
        for i,v in enumerate(f):
            h=False
            if isinstance(v,tuple): v,h=v
            cells[i].text=""; r=cells[i].paragraphs[0].add_run(v); r.font.size=Pt(fs)
            if h: r.font.highlight_color=YEL
    for row in tb.rows:
        for i,w in enumerate(anchos): row.cells[i].width=Cm(w)
    cap(t)

# ================================================================
H("2. ETAPA DE INFORMACIÓN",1)
P("En esta etapa se recoge toda la información necesaria para definir el producto antes de empezar a diseñarlo: costes, competencia, patentes, normativa, tecnologías disponibles y condiciones de almacenaje, embalaje y montaje. Con todo ello se redacta el documento EDP (Especificaciones de Diseño del Producto). Tal y como se indica en la metodología de la asignatura, esta etapa se mantiene abierta durante todo el proceso, y de hecho algunos valores de la etapa inicial se han corregido a partir de lo encontrado aquí.")

# ---------------- 2.1 COSTES ----------------
H("2.1. COSTES",2)
P("Para fijar el coste objetivo se ha partido del precio de venta de los productos de la competencia (apartado 2.2). Los taburetes plegables de dos peldaños con carga de 150 kg se venden en España entre unos 19 € (modelos básicos de aluminio) y 86 € (modelos de marca con peldaños anchos), con un precio medio de unos 54 €. El objetivo es situar el producto por debajo de las marcas de gama media-alta pero claramente por encima, en prestaciones, de los modelos más baratos.")
P("A partir de un precio de venta al público (PVP) objetivo de 45 € con IVA incluido se ha estimado, de forma aproximada, el coste de fabricación máximo que puede tener el producto:")
tabla(["Concepto","Cálculo","Importe"],[
("PVP objetivo (IVA incluido)","—","45,00 €"),
("PVP sin IVA","45 / 1,21","37,19 €"),
("Margen de la distribución (tienda)",("≈ 38 % del PVP sin IVA (hipótesis)",True),"− 14,13 €"),
("Precio de venta del fabricante","37,19 − 14,13","23,06 €"),
("Margen del fabricante",("≈ 25 % (hipótesis)",True),"− 5,77 €"),
("COSTE INDUSTRIAL OBJETIVO","","≈ 17,30 €")],[6,6,3.5],"Tabla 3. Cálculo del coste objetivo a partir del precio de venta. Elaboración propia.")
P("Este coste industrial se reparte, como primera aproximación, de la siguiente manera: materiales ≈ 8,50 €, transformación (corte, doblado, inyección, acabado) ≈ 4,00 €, montaje ≈ 1,50 €, embalaje ≈ 1,00 € y gastos generales y amortización de utillajes ≈ 2,30 €. Los porcentajes de margen son hipótesis habituales del sector y se revisarán cuando se tengan las masas reales de cada pieza a partir del modelo en SolidWorks.")
P("Para que estos costes sean alcanzables el producto tiene que fabricarse en series largas (del orden de 20.000 unidades al año), ya que, como se indica en el tema de EDP, solo las series largas permiten amortizar herramientas caras como los moldes de inyección o las matrices de estampación.")

# ---------------- 2.2 MERCADO ----------------
H("2.2. ESTUDIO DE MERCADO, PRODUCTOS SIMILARES Y SOLUCIONES",2)
P("Se han analizado los taburetes-escalera que se venden actualmente en las principales tiendas españolas (IKEA, Leroy Merlin, Obramat, ferreterías online). La Tabla 4 resume sus características principales y la Ilustración 5 compara sus precios.")
tabla(["Producto","Material","Peldaños / altura plataforma","Carga","Peso","Plegado","Precio"],[
("IKEA BEKVÄM","Haya maciza","2 / 50 cm (alto total)","100 kg","n. d.","No","24,99 €"),
("IKEA BOLMEN","PP (≥ 50 % reciclado)","1 / ≈ 25 cm","100 kg","n. d.","No","1,50 €"),
("Altipesa aluminio 2 peld.","Aluminio + plataforma de acero","2 / n. d.","150 kg","2 kg","Sí","19,27 €"),
("Hailo K20 (4396-901)","Acero","2 / 46 cm","150 kg","5,2 kg","Sí, 5 cm","58,87 €"),
("Hailo Safety (4312-001)","Acero, asidero 50 cm","2 / n. d.","150 kg","n. d.","Sí","52,39 €"),
("Hailo L90 (4442-701)","Aluminio","2 / ≈ 49 cm (alto total)","150 kg","n. d.","Sí, 19 cm","86,03 €"),
("Leroy Merlin, taburete doméstico","Acero pintado blanco","2 / 43 cm","150 kg","n. d.","Sí, 8 cm","n. d.")],
[3.6,3,2.6,1.4,1.3,1.6,1.6],"Tabla 4. Taburetes-escalera del mercado español (n. d.: dato no disponible). Fuentes: webs de los fabricantes y distribuidores, septiembre de 2026.",fs=8)
fig("figura5_precios_mercado.png","Ilustración 5. Precio de venta de los taburetes-escalera analizados. Elaboración propia.",15.5)
P("Análisis de las soluciones existentes:",bold=True)
B("Taburetes fijos de madera o plástico (IKEA BEKVÄM, BOLMEN): muy baratos y con buena estética, pero no se pliegan y su carga máxima es de solo 100 kg. El BEKVÄM incorpora un asa en el peldaño superior, una buena idea que se tendrá en cuenta.")
B("Taburetes plegables de aluminio básicos (Altipesa y similares): muy ligeros (2 kg) y baratos, pero el peldaño inferior tiene solo 80 mm de profundidad, por lo que no se apoya el pie entero y transmiten poca sensación de seguridad.")
B("Taburetes plegables de acero de marca (Hailo K20, Safety): muy robustos, peldaños anchos con alfombrilla antideslizante y plegado de solo 5 cm, pero pesan más de 5 kg y cuestan entre 52 y 59 €.")
B("Modelos de aluminio de gama alta (Hailo L90): peldaños anchos y ligeros, pero su precio (86 €) está fuera del alcance de muchos hogares y plegado ocupa 19 cm.")
P("Conclusiones del estudio de mercado:",bold=True)
P("Existe un hueco entre los modelos baratos pero poco cómodos y los modelos robustos pero pesados y caros. El nuevo diseño debe ofrecer peldaños profundos y la carga de 150 kg de los modelos de marca, con un peso y un precio más cercanos a los de los modelos básicos. Al comparar las propiedades fijadas en la etapa inicial con la competencia se ha visto que dos de ellas no eran suficientemente exigentes (el espesor plegado de 80 mm y el peso de 5 kg), así que, siguiendo la recomendación del tema de EDP de redefinir los requisitos cuando se detectan deficiencias frente a lo existente, se han cambiado a 50 mm y 3,5 kg respectivamente.")

# ---------------- 2.3 EDP ----------------
H("2.3. EDP (DOCUMENTO DE ESPECIFICACIONES DE DISEÑO DEL PRODUCTO)",2)
P("El EDP recoge las especificaciones y limitaciones que debe cumplir el producto, impuestas por el mercado y por la empresa, antes de empezar la fase creativa. Se ha redactado con frases cortas y sin conducir el diseño hacia una solución concreta, siguiendo los apartados vistos en la asignatura.")
edp=[("Almacenaje","Almacenable en interior durante al menos 2 años sin deterioro. Temperatura de 0 a 40 °C, humedad relativa ≤ 80 %. Cajas apilables sin dañar el producto."),
("Calidad y fiabilidad","Ningún fallo estructural con 150 kg durante toda su vida útil. Bloqueo fiable en todas las aperturas. Registro de fallos y devoluciones para retroalimentar el diseño."),
("Cantidad","Serie larga: unas 20.000 unidades/año. Justifica utillajes de inyección y estampación."),
("Competencia","Mejorar a los modelos básicos en profundidad de peldaño y seguridad percibida, y a los modelos de marca en peso y precio."),
("Consumidor / cliente","Hogares españoles y europeos. Usuario de cualquier edad adulta, con especial atención a personas mayores. Comprador en gran superficie de bricolaje, tienda de hogar u online."),
("Coste","PVP ≤ 45 € IVA incluido. Coste industrial ≤ 17,30 €."),
("Documentación","Manual con uso correcto e incorrecto, carga máxima, mantenimiento, reparación y reciclaje. Pictogramas de advertencia en el propio producto. Referencia a la UNE-EN 14183."),
("Embalaje","Caja de cartón reciclado, sin plásticos ni poliestireno. Debe proteger las esquinas y servir como expositor en tienda."),
("Entorno","Uso en interior (cocina, baño, salón, trastero). Resistente a la humedad del baño y a productos de limpieza domésticos."),
("Ergonomía","Alcance de 2,30 m para usuaria P5. Peldaños donde quepa el pie entero. Subida cómoda para personas mayores. Manejo y transporte con una mano."),
("Estética","Aspecto limpio y actual, colores neutros. Debe transmitir solidez."),
("Facilidad de fabricación","Pocas piezas, preferiblemente comunes a varias posiciones. Procesos estándar subcontratables."),
("Funcionalidad","Abrir y cerrar en ≤ 3 s. Bloqueo automático en posición de uso. Imposible de cerrar con el usuario encima."),
("Implicaciones sociopolíticas","Ninguna relevante. Favorece la autonomía de personas mayores en su hogar."),
("Instalación","No requiere instalación: se entrega montado y listo para usar."),
("Legislación","Reglamento (UE) 2023/988 de seguridad general de los productos. Responsabilidad por productos defectuosos (RDL 1/2007). REACH."),
("Mantenimiento","Sin mantenimiento periódico. Tacos y recubrimientos sustituibles como repuesto."),
("Normas","UNE-EN 14183 (taburetes-escalera) como norma de referencia."),
("Materiales","Reciclables, sin sustancias restringidas por REACH. Preferiblemente con contenido reciclado."),
("Patentes","No invadir patentes en vigor en los países de venta (ver apartado 2.5)."),
("Peso","≤ 3,5 kg."),
("Planificación","Diseño y prototipo en 6 meses; ensayos y preserie en 3 meses más."),
("Procesos","Definir los procesos de cada componente y los tratamientos superficiales (pintura, anodizado…) en la etapa de diseño de detalle."),
("Reciclaje","≥ 90 % del peso reciclable. Piezas de un solo material, separables con herramientas comunes."),
("Residuos","Minimizar recortes y mermas de fabricación. Sin piezas que acaben en vertedero por imposibles de separar."),
("Seguridad","Carga máxima 150 kg. Sin aristas vivas. Holguras que eviten atrapar los dedos. Antideslizante en peldaños y apoyos."),
("Restricciones de empresa","Aprovechar procesos subcontratados (tubo, inyección, estampación). Montaje final en la propia empresa."),
("Restricciones de mercado","Producto para el mercado europeo: marcas y textos en varios idiomas."),
("Tamaño y volumen","Abierto: huella ≈ 500 × 450 mm. Plegado: espesor ≤ 50 mm."),
("Transporte","Por carretera en palé europeo. Caja con asa o hendidura para cogerla."),
("Vida del producto (mercado)","Permanencia estimada en el mercado de 8 a 10 años con pequeñas actualizaciones de color y acabado."),
("Vida útil","≥ 10 años en uso doméstico.")]
tabla(["Especificación","Requisitos y limitaciones"],edp,[4,12],"Tabla 5. Documento EDP del taburete-escalera. Elaboración propia.")

# ---------------- 2.4 TECNOLOGÍAS ----------------
H("2.4. TECNOLOGÍAS EXISTENTES",2)
P("A partir de los productos analizados se han identificado las tecnologías de fabricación que se usan actualmente en este tipo de productos:")
tabla(["Componente","Tecnología","Ventajas","Inconvenientes"],[
("Bastidor / largueros","Tubo de acero curvado y soldado, pintura en polvo","Muy resistente, barato, fácil de fabricar","Pesado; la pintura puede saltar"),
("Bastidor / largueros","Perfil de aluminio extruido, unido con remaches","Ligero, no se oxida, perfil a medida","Material más caro y con más energía de producción primaria"),
("Bastidor / largueros","Madera maciza o contrachapado","Estética cálida, material renovable","Mayor espesor, sensible a la humedad"),
("Peldaños","Chapa de acero estampada con relieve","Resistente y barata en series largas","Peso; necesita recubrimiento antideslizante"),
("Peldaños","Aluminio extruido estriado","Ligero, antideslizante por su propia forma","Profundidad limitada por el perfil"),
("Peldaños","Polipropileno reforzado con fibra de vidrio inyectado","Formas complejas, nervios, antideslizante integrado","Molde caro, solo rentable en series largas"),
("Uniones","Remaches, pasadores y tornillería","Desmontables (tornillos) o baratos (remaches)","Los remaches dificultan el desmontaje"),
("Apoyos","Tacos de caucho o TPE (elastómero termoplástico)","Antideslizantes, no marcan el suelo","Desgaste; deben poder sustituirse"),
("Acabados","Pintura en polvo, anodizado, lacado","Protección y color","Impacto ambiental del tratamiento")],
[2.8,4.2,4.4,4.4],"Tabla 6. Tecnologías existentes para la fabricación de taburetes-escalera. Elaboración propia.")

# ---------------- 2.5 PATENTES ----------------
H("2.5. ESTUDIO DE PATENTES",2)
P("Antes de diseñar un producto para venderlo es necesario conocer las patentes que existen sobre productos similares, ya que podrían impedir su comercialización. Una patente es un título que da a su titular el derecho exclusivo de explotar una invención durante 20 años en el territorio en el que se ha concedido, a cambio de hacerla pública. En España existe además el modelo de utilidad, que protege invenciones de menor rango (una configuración o estructura que da una ventaja práctica) durante 10 años. Hay que tener en cuenta que la protección es territorial: una patente de Estados Unidos no tiene efectos en España.")
P("La búsqueda se ha hecho en Google Patents, Espacenet (Oficina Europea de Patentes) e INVENES (Oficina Española de Patentes y Marcas), con palabras clave como «taburete plegable», «escalera plegable», «folding step stool» y «Klapptritt». A continuación se analizan las cuatro más cercanas a nuestro producto.")
pat=[("2.5.1. PATENTE PRIMERA: ES 1 070 877 U — Dispositivo de enclavamiento para taburetes plegables y similares",
  [("Tipo: ",False),("modelo de utilidad español. ",False),("Solicitante: ",False),("Gesco, S.L. ",False),("Inventor: ",False),("Jorge Calvo Pajares. ",False),("Nº de solicitud: ",False),("U200901059, presentada el 26/06/2009. ",False),("Estado: ",False),("caducado.",False)],
  "Resumen: dispositivo de enclavamiento para taburetes plegables destinado a mantener el taburete en la posición de desplegado con plenas garantías de seguridad.",
  [("Reivindicación principal: dispositivo de enclavamiento que, montado entre los elementos articulados del taburete, bloquea automáticamente la estructura al desplegarla e impide su plegado mientras no se accione el propio dispositivo. ",False),("[Comprobar y copiar el texto literal de la reivindicación 1 en Google Patents / INVENES]",True)],
  "¿Afecta a nuestro producto? No. Aunque nuestro taburete también llevará un bloqueo en posición abierta, el modelo de utilidad está caducado (su protección máxima era de 10 años, hasta 2019), por lo que su contenido es de dominio público y puede usarse libremente."),
("2.5.2. PATENTE SEGUNDA: US 6 454 050 B2 — Foldable step stool with leg lock and handle",
  [("Tipo: ",False),("patente de Estados Unidos. ",False),("Titular original: ",False),("Cosco Management Inc. (actualmente Dorel Home Furnishings). ",False),("Fecha de solicitud: ",False),("29/06/2001, con prioridad de 11/08/2000. ",False),("Estado: ",False),("expirada.",False)],
  "Resumen: taburete plegable con un bastidor formado por una pata delantera y una pata trasera que se mueven una respecto de la otra, y un asa de transporte que gira sobre la pata delantera. Al girar el asa, una pieza de retención unida a ella bloquea la pata trasera con la delantera.",
  [("Reivindicación principal: taburete que comprende un bastidor con pata delantera y pata trasera acopladas para moverse entre sí, un asa, un soporte de pivote que permite el giro del asa sobre la pata delantera y un retenedor unido al asa que atrapa una parte de la pata trasera para bloquear ambas patas cuando el asa llega a una posición determinada.",False)],
  "¿Afecta a nuestro producto? No. Es una patente estadounidense, sin efectos en España, y además ha expirado (más de 20 años desde su solicitud). Es interesante porque combina asa y bloqueo en una sola pieza, una idea que podría aprovecharse."),
("2.5.3. PATENTE TERCERA: US 6 966 404 B2 — Folding step stool",
  [("Tipo: ",False),("patente de Estados Unidos. ",False),("Inventor: ",False),("Paul K. Meeker. ",False),("Titular original: ",False),("Cosco Management Inc. ",False),("Fecha de solicitud: ",False),("19/05/2003. ",False),("Estado: ",False),("expirada (impago de tasas).",False)],
  "Resumen: taburete plegable con un escalón superior y al menos dos patas plegables unidas a él. Unos mecanismos de bloqueo, que deslizan por la cara inferior del escalón y giran respecto a cada pata, mantienen las patas en posición abierta; al accionarlos, las patas se pliegan paralelas al escalón. El escalón tiene unas aberturas que forman un asa de transporte.",
  [("Reivindicación principal: taburete que comprende un escalón, al menos dos patas plegables unidas al escalón y un mecanismo de bloqueo acoplado de forma deslizante a la cara inferior del escalón y de forma pivotante a cada pata, que bloquea y sostiene cada pata en la posición abierta de uso.",False)],
  "¿Afecta a nuestro producto? No, por los mismos motivos que la anterior (patente estadounidense y expirada). Su idea del asa formada por aberturas en el escalón es aplicable a la plataforma de nuestro taburete."),
("2.5.4. PATENTE CUARTA: US 12 419 421 B2 — Foldable locking step stool",
  [("Tipo: ",False),("patente de Estados Unidos. ",False),("Titular: ",False),("Jool Products LLC. ",False),("Inventores: ",False),("Judah Bergman, Yishak Sultan y otros. ",False),("Nº de solicitud: ",False),("US 17/743,098 (publicada como US 2023/0363538 A1). ",False),("Estado: ",False),("en vigor (concedida en 2025).",False)],
  "Resumen: taburete plegable cuya plataforma se divide en dos mitades unidas por una bisagra. Dispone de un mecanismo de bloqueo a prueba de niños formado por dos deslizadores que hay que accionar a la vez para liberar unos pestillos y poder plegarlo, evitando así que el usuario se pille los dedos. Incluye agarres sobremoldeados y un asa integrada que queda oculta durante el uso.",
  [("Reivindicación principal: taburete plegable que comprende una plataforma que se pliega mediante bisagra en un primer y un segundo miembro, patas unidas con bisagra a la plataforma y un mecanismo de bloqueo a prueba de niños, de forma que en posición abierta ambos miembros quedan bloqueados en un mismo plano y en posición cerrada quedan desbloqueados y plegados; el mecanismo comprende un par de elementos deslizantes que deben accionarse simultáneamente para desbloquearlo.",False)],
  [("¿Afecta a nuestro producto? En principio no. La patente solo tiene efectos en Estados Unidos y nuestro diseño no divide la plataforma en dos mitades ni usa dos deslizadores simultáneos. ",False),("[Comprobar en Espacenet si existe una solicitud europea (EP) o internacional (WO) de la misma familia]",True)])]
for tit,datos,res,rei,afe in pat:
    d.add_heading(tit,3)
    P(*datos)
    P(res)
    P(*rei)
    P(("[Insertar aquí el dibujo principal de la patente]",True))
    if isinstance(afe,list): P(*afe)
    else: P(afe)
P("Otras patentes consultadas: EP 2 660 419 A2 (Hailo Werk, «Steiggerät, nämlich Klapptritt oder Trittleiter»), EP 3 056 656 A1 (Hailo Werk, soporte para pies de escaleras y taburetes) y US 6 390 238 B1 (Cosco, versión anterior de la patente segunda).")
P("¿Es patentable nuestro producto?",bold=True)
P("El concepto general de taburete-escalera plegable de dos peldaños no es patentable, porque no cumple el requisito de novedad: existen desde hace décadas y hay muchas patentes sobre él, como las analizadas. Sin embargo, sí podría protegerse una solución concreta que se desarrolle durante el diseño, por ejemplo un mecanismo de bloqueo que combine el asa de transporte con el cierre de seguridad de una forma nueva. Al tratarse de una mejora en la configuración o estructura del producto que da una ventaja práctica, la vía más adecuada sería un modelo de utilidad en España, más rápido y barato que una patente. Además, la apariencia externa del producto podría protegerse como diseño industrial y el nombre comercial como marca.")

# ---------------- 2.6 LEGISLACIÓN ----------------
H("2.6. LEGISLACIÓN A CUMPLIR",2)
P("Los taburetes-escalera no tienen una directiva europea propia (no llevan marcado CE), por lo que deben cumplir la legislación general de seguridad de los productos de consumo. La norma UNE-EN 14183 se usa como referencia técnica para demostrar que el producto es seguro.")
tabla(["Documento","Contenido y aplicación al producto","Carácter"],[
("Reglamento (UE) 2023/988, de seguridad general de los productos","Aplicable desde el 13/12/2024; deroga la Directiva 2001/95/CE. Obliga a que todo producto de consumo sea seguro, a hacer una evaluación de riesgos documentada y conservarla 10 años, y a identificar al fabricante en el producto.","Obligatorio"),
("Real Decreto Legislativo 1/2007 (Ley General para la Defensa de los Consumidores y Usuarios)","Libro III: responsabilidad civil por daños causados por productos defectuosos (defectos de diseño, de fabricación o de información).","Obligatorio"),
("Reglamento (CE) 1907/2006 (REACH)",("Restringe sustancias peligrosas en los artículos. Afecta a tacos y recubrimientos de caucho o TPE en contacto con la piel (anexo XVII, entrada 50: HAP ≤ 1 mg/kg).",True),"Obligatorio"),
("Ley 7/2022 de residuos y suelos contaminados y RD 1055/2022 de envases","Gestión de residuos y responsabilidad ampliada del productor sobre los envases (caja de cartón).","Obligatorio"),
(("Reglamento (UE) 2025/40 de envases y residuos de envases",True),"Reducción de envases, contenido reciclable y etiquetado de los materiales del embalaje.","Obligatorio (aplicación progresiva)"),
("UNE-EN 14183 Taburetes-escalera","Taburetes de hasta 1 m de altura de plataforma y carga máxima de 150 kg. Define diseño, dimensiones, materiales, ensayos de resistencia y de deslizamiento lateral, y marcado. Peldaños a distancias iguales. Barandilla obligatoria si la plataforma es menor de 24 × 40 cm y está a más de 75 cm.","Voluntaria (referencia)"),
("UNE-EN ISO 7250-1","Definiciones de las medidas antropométricas básicas del cuerpo humano.","Voluntaria"),
("UNE-EN ISO 14006 y UNE-EN ISO 14040/14044","Integración del ecodiseño en el proceso de diseño y metodología del análisis de ciclo de vida.","Voluntaria")],
[4.6,9,2.4],"Tabla 7. Legislación y normativa aplicable. Elaboración propia.")
P("Cabe señalar que el comité europeo de normalización está preparando la norma prEN 131-9, dedicada a los taburetes-escalera dentro de la familia de normas de escaleras, que en el futuro podría sustituir a la EN 14183.")

# ---------------- 2.7 NORMAS OTROS PAÍSES ----------------
H("2.7. NORMAS INTERESANTES DE OTROS PAÍSES",2)
B("ANSI ASC A14.11-2018 (Estados Unidos): norma específica de seguridad para taburetes-escalera de madera, metal, plástico y plástico reforzado. Define cargas de servicio de 200, 225, 250, 300 y 375 libras (unos 90 a 170 kg). Antes de 2018 estos productos se regulaban dentro de las normas de escaleras A14.1, A14.2 y A14.5.")
B("DIN EN 14183 (Alemania) y BS EN 14183 (Reino Unido): versiones nacionales de la norma europea. En Alemania es habitual además la marca voluntaria GS («Geprüfte Sicherheit», seguridad comprobada), muy valorada por el consumidor.")
B("EN 131 (Europa): norma de escaleras, que excluye expresamente a los taburetes-escalera, pero que algunos fabricantes también cumplen para dar más confianza al comprador (como el modelo de Leroy Merlin analizado, certificado según EN 131 y EN 14183).")
P("Si el producto se quisiera vender en Estados Unidos habría que diseñarlo para una carga de servicio de al menos 300 libras (136 kg), algo que ya se cumple con los 150 kg fijados.")

# ---------------- 2.8 ENTORNO ----------------
H("2.8. AFECCIÓN AL ENTORNO",2)
P("Siguiendo el enfoque de ecodiseño «de la cuna a la tumba», se han identificado de forma cualitativa las entradas, salidas e impactos del producto en cada fase de su ciclo de vida. El análisis cuantitativo (ACV) se hará en la etapa de evaluación, comparando las alternativas de material.")
tabla(["Fase","Entradas","Salidas / impactos","Medidas de ecodiseño"],[
("Materias primas y fabricación","Acero o aluminio, plásticos, energía","Emisiones de CO₂ (sobre todo en el aluminio primario), recortes, COV de la pintura","Material reciclado, menos material, pintura en polvo sin disolventes"),
("Distribución","Combustible, embalaje","Emisiones del transporte, residuos de embalaje","Plegado plano: más unidades por palé; embalaje solo de cartón"),
("Uso","Ninguna (no consume energía)","Desgaste de tacos","Tacos sustituibles, vida útil larga"),
("Fin de vida","Transporte a punto limpio","Chatarra metálica, plástico","Piezas monomaterial y desmontables; ≥ 90 % reciclable")],
[3,3.4,4.6,5],"Tabla 8. Afección al entorno a lo largo del ciclo de vida. Elaboración propia.")
P("En la Ilustración 6 se compara, mediante la rueda estratégica de ecodiseño vista en la asignatura, un taburete de acero típico del mercado con el objetivo del nuevo diseño. Las mejoras se centran en la selección de materiales, la reducción de material, la distribución (gracias al plegado plano) y el fin de vida.")
fig("figura6_rueda_ecodiseno.png","Ilustración 6. Rueda estratégica de ecodiseño: producto actual frente a objetivo. Valoración cualitativa propia.",12)

# ---------------- 2.9 ALMACENAJE ----------------
H("2.9. ALMACENAJE",2)
P("Desde que se fabrica hasta que llega al usuario, el taburete puede pasar varios meses en almacenes. Se establecen las siguientes condiciones:")
B("Almacenaje en interior, protegido de la lluvia y de la luz solar directa, que podría degradar los tacos y las piezas de plástico.")
B("Temperatura entre 0 y 40 °C y humedad relativa inferior al 80 %, para evitar la corrosión de las partes de acero.")
B("Durabilidad almacenado: al menos 2 años sin pérdida de propiedades.")
B("Cajas apiladas en el palé en un máximo de 3 alturas, de canto, para no aplastar el cartón.")
P("En casa del usuario, el producto está pensado para guardarse plegado en huecos estrechos (detrás de una puerta, entre un mueble y la pared) o colgado en la pared mediante su asa.")

# ---------------- 2.10 EMBALAJE ----------------
H("2.10. EMBALAJE",2)
P("El embalaje tiene tres funciones: proteger el producto durante el transporte, permitir su almacenaje y servir como reclamo publicitario en la tienda. Se proponen las siguientes características:")
B("Caja de cartón ondulado de canal simple, fabricada con cartón reciclado y totalmente reciclable. Sin bolsas de plástico ni piezas de poliestireno; las esquinas se protegen con cantoneras de cartón.")
B("Dimensiones aproximadas de la caja, a partir del tamaño plegado: 640 × 470 × 60 mm.")
B("Impresión a una tinta con la foto del producto, la carga máxima, la norma que cumple y los pictogramas de uso.")
P("Paletización: en un palé europeo (1200 × 800 mm), colocando las cajas de canto se pueden poner 20 cajas por fila en el lado de 1200 mm y 3 alturas de 470 mm, lo que da 60 unidades por palé con una altura total de unos 1,55 m. Si las cajas se colocaran planas cabrían solo unas 50 unidades, por lo que se elige la disposición de canto. El plegado plano del producto es la clave para reducir el volumen transportado y, con ello, el coste y las emisiones del transporte.")

# ---------------- 2.11 MONTAJE ----------------
H("2.11. MONTAJE",2)
P("El producto se entrega completamente montado, por lo que el usuario no tiene que realizar ninguna operación salvo desplegarlo. El montaje se hace en fábrica, con una secuencia prevista como la siguiente:")
for t in ["Preparación de los largueros delanteros y traseros (cortados, taladrados y con su acabado superficial).",
"Colocación de los tacos antideslizantes en los extremos de los largueros.",
"Unión de los peldaños y de la plataforma a los largueros delanteros mediante pasadores o remaches.",
"Unión de la pata trasera a la plataforma en la articulación superior.",
"Montaje de los tirantes y del mecanismo de bloqueo.",
"Colocación de las etiquetas de advertencia y carga máxima.",
"Control final: apertura y cierre, comprobación del bloqueo y prueba de carga por muestreo."]:
    N(t)
P("Se procurará usar el menor número posible de piezas distintas y de tipos de unión, y que todas las uniones puedan deshacerse con herramientas comunes para facilitar la reparación y el reciclaje.")

# ================================================================
H("3. ETAPA TÉCNICO-CREATIVA",1)
P("El objetivo de esta etapa es buscar soluciones que cumplan los objetivos y el EDP. Primero se ha usado una técnica de pensamiento divergente (mapa mental) para generar ideas sin criticarlas, y después el método TRIZ para resolver las contradicciones técnicas del producto.")
H("3.1. MAPA MENTAL",2)
P("En el mapa mental se han ordenado, alrededor del producto, todas las ideas relacionadas con el usuario, la seguridad, el almacenaje, los materiales, la fabricación, la estética, el medio ambiente y el coste. Sirve para no olvidar ningún aspecto al generar las propuestas.")
fig("figura7_mapa_mental.png","Ilustración 7. Mapa mental del taburete-escalera. Elaboración propia.",16)

H("3.2. MÉTODO TRIZ",2)
P("TRIZ (Teoría para Resolver Problemas de Inventiva) fue desarrollado por Genrich Altshuller a partir del análisis de miles de patentes. Altshuller observó que la mayoría de los problemas de diseño aparecen cuando se intenta mejorar un parámetro del producto y, al hacerlo, empeora otro (contradicción técnica). Para resolverlas definió 39 parámetros técnicos y 40 principios inventivos, relacionados entre sí mediante la matriz de contradicciones: en la fila se busca el parámetro que se quiere mejorar, en la columna el que empeora, y la celda indica los principios que más veces se han usado para resolver ese conflicto.")
par=["Peso del objeto móvil","Peso del objeto inmóvil","Longitud del objeto móvil","Longitud del objeto inmóvil","Área del objeto móvil","Área del objeto inmóvil","Volumen del objeto móvil","Volumen del objeto inmóvil","Velocidad","Fuerza","Tensión o presión","Forma","Estabilidad del objeto","Resistencia","Durabilidad del objeto móvil","Durabilidad del objeto inmóvil","Temperatura","Brillo","Energía gastada por el objeto móvil","Energía gastada por el objeto inmóvil","Potencia","Pérdida de energía","Pérdida de sustancia","Pérdida de información","Pérdida de tiempo","Cantidad de sustancia","Fiabilidad","Precisión de medida","Precisión de fabricación","Factores dañinos que actúan sobre el objeto","Factores dañinos generados por el objeto","Facilidad de fabricación","Facilidad de uso","Facilidad de reparación","Adaptabilidad","Complejidad del dispositivo","Complejidad de control","Nivel de automatización","Productividad"]
pri=["Segmentación","Extracción","Calidad local","Asimetría","Combinación","Universalidad","Anidamiento","Contrapeso","Acción previa contraria","Acción previa","Amortiguación anticipada","Equipotencialidad","Inversión","Curvatura","Dinamicidad","Acción parcial o excesiva","Otra dimensión","Vibración mecánica","Acción periódica","Continuidad de la acción útil","Acción rápida","Convertir lo dañino en útil","Retroalimentación","Intermediario","Autoservicio","Copia","Objetos baratos de vida corta","Sustitución de sistemas mecánicos","Neumática e hidráulica","Membranas flexibles","Materiales porosos","Cambio de color","Homogeneidad","Descarte y recuperación","Cambio de parámetros","Cambio de fase","Expansión térmica","Oxidantes fuertes","Atmósfera inerte","Materiales compuestos"]
filas=[]
for i in range(40):
    filas.append((str(i+1) if i<39 else "",par[i] if i<39 else "",str(i+1),pri[i]))
tabla(["Nº","Parámetro técnico","Nº","Principio inventivo"],filas,[1,6.5,1,6.5],"Tabla 9. Los 39 parámetros técnicos y los 40 principios inventivos de Altshuller.",fs=8)

d.add_heading("Contradicción primera: resistencia frente a peso",3)
P("Para soportar 150 kg el taburete tiene que ser resistente, pero si se refuerza la estructura (más material, perfiles más gruesos) aumenta su peso, y se ha fijado un máximo de 3,5 kg para poder llevarlo con una mano.")
B("Parámetro que se quiere mejorar: 14 – Resistencia.")
B("Parámetro que empeora: 1 – Peso del objeto móvil.")
B("Principios propuestos por la matriz: 1 (segmentación), 8 (contrapeso), 15 (dinamicidad) y 40 (materiales compuestos).")
P("Aplicación de los principios al producto:")
B("Principio 1, segmentación: dividir la estructura en piezas independientes, cada una del material y la sección justos para su función (largueros de perfil resistente, peldaños nervados), en lugar de piezas macizas. Además facilita la sustitución de piezas y el reciclaje.")
B("Principio 15, dinamicidad: que los peldaños y la plataforma giren con el larguero y se coloquen en la posición óptima en cada momento: horizontales para trabajar y verticales para guardar.")
B("Principio 40, materiales compuestos: usar peldaños de polipropileno reforzado con fibra de vidrio, con nervios interiores, que dan mucha rigidez con poco peso.")
B("Principio 8, contrapeso: no se ve una aplicación directa, aunque se aprovecha la idea de que sea el propio peso del usuario el que mantenga el bloqueo cerrado (ver contradicción tercera).")

d.add_heading("Contradicción segunda (física): base grande y base pequeña",3)
P("Para ser estable, el taburete necesita una base de apoyo amplia, de modo que el centro de gravedad del usuario quede siempre dentro de ella. Pero para guardarlo, el producto tiene que ser lo más pequeño y plano posible. Es decir, el mismo elemento tiene que ser grande y pequeño a la vez, lo que en TRIZ se llama una contradicción física.")
P("Este tipo de contradicción se resuelve con los principios de separación. En este caso se aplica la separación en el tiempo: la base es grande cuando se usa y pequeña cuando se guarda. Esto lleva directamente al plegado (principio 15, dinamicidad) y a alojar los peldaños dentro del hueco de los largueros al cerrarlo (principio 7, anidamiento), para conseguir el espesor plegado de 50 mm.")

d.add_heading("Contradicción tercera: facilidad de uso frente a fiabilidad",3)
P("Se busca que el taburete se abra y se cierre muy rápido y con una sola mano, pero cuanto más fácil sea cerrarlo, mayor es el riesgo de que se cierre accidentalmente con el usuario encima.")
B("Parámetro que se quiere mejorar: 33 – Facilidad de uso.")
B("Parámetro que empeora: 27 – Fiabilidad.")
P(("Principios propuestos por la matriz: [consultar la celda 33 / 27 en el software TRIZ40 y completar]",True))
P("Independientemente de la celda, para este problema se consideran especialmente útiles los siguientes principios:")
B("Principio 10, acción previa: que el bloqueo se active solo al abrir el taburete, sin que el usuario tenga que hacer nada.")
B("Principio 11, amortiguación anticipada: añadir un segundo seguro que haya que accionar para plegar (dos gestos), como hace la patente cuarta con sus dos deslizadores.")
B("Principio 22, convertir lo dañino en útil: diseñar el tirante para que quede pasado de su punto muerto, de modo que el peso del usuario, en lugar de tender a cerrar el taburete, lo mantenga más bloqueado.")

H("3.3. BOCETO INICIAL",2)
P("A partir del mapa mental y de las soluciones de TRIZ se ha dibujado el boceto inicial del concepto (Ilustración 8). Se trata de un taburete de tipo tijera con dos peldaños iguales de 240 mm de altura. Los largueros delanteros llevan el peldaño inferior y la plataforma superior, y la pata trasera se articula bajo la plataforma. Un tirante une el peldaño inferior con la pata trasera y bloquea la apertura. Al plegarlo, el peldaño y la plataforma giran hasta quedar verticales, alojados entre los largueros, y todo el conjunto queda en un plano de unos 50 mm de espesor.")
fig("figura8_boceto_inicial.png","Ilustración 8. Boceto inicial del concepto. Elaboración propia.",16.5)
P("Este mismo concepto se desarrollará con tres alternativas de material, que se compararán en la etapa de evaluación mediante un análisis de ciclo de vida:")
B("Alternativa 1: largueros de tubo rectangular de acero S235 con pintura en polvo y peldaños de chapa de acero estampada.")
B("Alternativa 2: largueros de perfil extruido de aluminio 6063 y peldaños de polipropileno reforzado con fibra de vidrio.")
B("Alternativa 3: largueros y peldaños de contrachapado de abedul, con herrajes y tirantes de acero.")

H("3.4. DISEÑO EN SOLID DEL PRODUCTO FINAL",2)
P(("[Pendiente: modelado en SolidWorks del conjunto, con una configuración por cada alternativa de material, y capturas del modelo]",True))

# ---------------- Bibliografía ----------------
d.add_heading("Bibliografía",2)
bib+=["IKEA España. BEKVÄM taburete escalón, haya. https://www.ikea.com/es/es/p/bekvam-taburete-con-escalon-haya-60178887/",
"IKEA. BOLMEN taburete escalón. https://www.ikea.com/es/es/p/bolmen-taburete-escalon-blanco-70574429/",
"Obramat. Taburete de aluminio 2 peldaños. https://www.obramat.es/productos/taburete-de-aluminio-2-peldanos-10455151.html",
"ManoMano. Hailo 4396-901 Taburete de acero K20 (2 peldaños). https://www.manomano.es/p/hailo-4396-901---taburete-de-acero-k20-2-peldaos-37584210",
"Comercial Pazos. Hailo 4442-701 Taburete de aluminio L90. https://comercialpazos.com/hailo-4442-701-taburete-de-aluminio-l90-stepke-2x2-peldanos",
"Leroy Merlin. Taburete plegable doméstico de acero - 2 peldaños. https://www.leroymerlin.es/productos/taburete-plegable-domestico-de-acero-2-peldanos-88754256.html",
"Google Patents. ES1070877U, US6454050B2, US6966404B2, US12419421B2. https://patents.google.com",
"Test and Research Centre. Step Ladders vs. Step Stools: understanding EN 131 and EN 14183. https://testandresearch.org/step-ladders-vs-step-stools-understanding-en-131-and-en-14183/",
"AFESPO. EN 14183, normativa obligatoria de los taburetes. https://www.afespo.com/en-14183-normativa-obligatoria-de-taburetes/",
"Reglamento (UE) 2023/988 del Parlamento Europeo y del Consejo, relativo a la seguridad general de los productos. https://eur-lex.europa.eu/legal-content/es/ALL/?uri=CELEX:32023R0988",
"ANSI Blog. New Standard for Stepstools – ANSI ASC A14.11-2018. https://blog.ansi.org/ansi/standard-stepstools-ansi-asc-a14-11-2018/",
"Altshuller, G. Matriz de contradicciones y 40 principios inventivos. Herramienta TRIZ40: https://www.triz40.com",
"Sanz Adán, F. Apuntes de Ingeniería Simultánea: Metodologías, EDP, Patentes, Ecodiseño, Técnicas de evaluación. Universidad de La Rioja."]
for t in bib: N(t)
d.save("Trabajo_Taburete.docx")
print("ok",len(d.inline_shapes),"imgs",len(d.tables),"tablas")
