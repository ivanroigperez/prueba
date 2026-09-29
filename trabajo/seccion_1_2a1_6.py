from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
d=Document("Trabajo_Taburete.docx")
# quitar la bibliografía del final para volver a ponerla al final
body=d.element.body
idx=[i for i,p in enumerate(d.paragraphs) if p.text.startswith("Bibliografía")][0]
bib=[p for p in d.paragraphs[idx:]]
bib_txt=[p.text for p in bib[1:]]
for p in bib: p._p.getparent().remove(p._p)

P=d.add_paragraph
def H(t,l): d.add_heading(t,l)
def fig(path,txt,w=15):
    p=P(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(path,width=Cm(w))
    c=P(); c.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=c.add_run(txt); r.italic=True; r.font.size=Pt(9)
def shade(cell,color):
    tcPr=cell._tc.get_or_add_tcPr(); s=OxmlElement("w:shd"); s.set(qn("w:val"),"clear"); s.set(qn("w:color"),"auto"); s.set(qn("w:fill"),color); tcPr.append(s)
def tabla(cab,filas,anchos,cap):
    t=d.add_table(rows=1,cols=len(cab)); t.style="Table Grid"
    for i,c in enumerate(cab):
        cell=t.rows[0].cells[i]; cell.text=""; r=cell.paragraphs[0].add_run(c); r.bold=True; r.font.size=Pt(9.5); r.font.color.rgb=RGBColor(255,255,255); shade(cell,"C0392B")
    for f in filas:
        cells=t.add_row().cells
        for i,v in enumerate(f):
            cells[i].text=""; r=cells[i].paragraphs[0].add_run(v); r.font.size=Pt(9)
    for row in t.rows:
        for i,w in enumerate(anchos): row.cells[i].width=Cm(w)
    c=P(); c.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=c.add_run(cap); r.italic=True; r.font.size=Pt(9)

# ---------- 1.2 ----------
H("1.2. DEFINICIÓN",2)
P("El producto a diseñar es un taburete-escalera plegable de uso doméstico, es decir, un elemento auxiliar de pocos peldaños que permite a una persona elevarse de forma segura para acceder a zonas altas de la vivienda (muebles altos de cocina, altillos, estanterías, lámparas, cortinas…), y que cuando no se utiliza se puede plegar para guardarlo en poco espacio.")
P("No se trata de una escalera de mano ni de una escalera de tijera profesional, sino de un producto pensado para un uso breve y frecuente dentro de casa, por cualquier miembro de la familia.")

# ---------- 1.3 ----------
H("1.3. OBJETIVOS",2)
P("Objetivo general:").runs[0].bold=True
P("Diseñar un taburete-escalera doméstico que sea seguro y estable durante su uso y que, a la vez, sea ligero y fácil de guardar, con un precio competitivo y un impacto ambiental reducido a lo largo de su ciclo de vida.")
P("Objetivos específicos:").runs[0].bold=True
for t in ["Permitir que una usuaria de percentil 5 alcance con comodidad una altura de 2,30 m, lo que cubre el mueble alto completo y la balda media del altillo.",
"Soportar a un usuario adulto de hasta 150 kg sin riesgo de rotura, deformación ni vuelco.",
"Conseguir que, una vez plegado, ocupe el mínimo espesor posible para guardarlo en huecos estrechos.",
"Que se pueda abrir, cerrar y transportar con una sola mano.",
"Cumplir la norma UNE-EN 14183 sobre taburetes-escalera.",
"Reducir la huella de carbono del producto respecto a los modelos equivalentes del mercado, utilizando materiales reciclables y fáciles de separar.",
"Mantener un precio de venta similar o inferior al de la competencia de gama media."]:
    d.add_paragraph(t,style="List Number")
P("Para fijar la altura necesaria se ha hecho un cálculo sencillo a partir de los datos del apartado anterior. Si el objetivo es llegar a 2,30 m y el alcance de una usuaria de percentil 5 es de unos 1,85 m, la plataforma superior tiene que estar como mínimo a 0,45 m del suelo. Por comodidad se toma una altura de 0,48 m, repartida en dos peldaños de unos 0,24 m (Ilustración 3).")
fig("figura3_altura_plataforma.png","Ilustración 3. Determinación de la altura de la plataforma superior. Elaboración propia.",15)

# ---------- 1.4 ----------
H("1.4. FUNCIONES",2)
P("Siguiendo la clasificación vista en la asignatura, las funciones del producto se dividen en funciones de uso, de manipulación y comunicativas. La función principal es elevar al usuario de forma segura para que pueda alcanzar zonas altas; el resto son funciones secundarias que hacen posible o mejoran la principal (Ilustración 4).")
fig("figura4_arbol_funciones.png","Ilustración 4. Árbol de funciones del taburete-escalera. Elaboración propia.",15.5)
P("Funciones de uso (principales):").runs[0].bold=True
for t in ["Soportar el peso del usuario mientras sube, está de pie y baja.",
"Mantener la estabilidad aunque el usuario se incline o cambie el peso de un pie a otro.",
"Proporcionar un apoyo firme y cómodo para el pie.",
"Ofrecer un punto de agarre en el que apoyarse al subir y bajar."]:
    d.add_paragraph(t,style="List Bullet")
P("Funciones de manipulación (secundarias):").runs[0].bold=True
for t in ["Desplegarse y plegarse de forma rápida.",
"Quedar bloqueado en posición de uso para que no se cierre solo.",
"Poder transportarse de una habitación a otra con una mano.",
"Almacenarse en poco espacio cuando no se usa."]:
    d.add_paragraph(t,style="List Bullet")
P("Funciones comunicativas (secundarias):").runs[0].bold=True
for t in ["Transmitir sensación de seguridad y solidez al usuario.",
"Indicar la carga máxima admisible y las advertencias de uso.",
"Dejar claro, a simple vista, cómo se abre y cuándo está bien bloqueado.",
"Integrarse estéticamente en la cocina o en el salón, ya que muchas veces estará a la vista."]:
    d.add_paragraph(t,style="List Bullet")

# ---------- 1.5 ----------
H("1.5. REQUISITOS",2)
P("Los requisitos son las condiciones que tiene que cumplir el producto para realizar correctamente sus funciones. Se han agrupado en los tipos vistos en clase (ergonómicos, estéticos, de uso y medioambientales), añadiendo los de seguridad y los económicos, y se han separado en imprescindibles (I) y deseables (D).")
req=[("R1","Seguridad","Soportar a un adulto de hasta 150 kg","I"),
("R2","Seguridad","No volcar al subir, bajar o inclinarse hacia un lado","I"),
("R3","Seguridad","No poder plegarse accidentalmente mientras se usa","I"),
("R4","Seguridad","Peldaños y apoyos antideslizantes","I"),
("R5","Seguridad","Cumplir la norma UNE-EN 14183","I"),
("R6","Ergonómico","Permitir alcanzar 2,30 m a una usuaria de percentil 5","I"),
("R7","Ergonómico","Peldaños en los que quepa el pie completo","I"),
("R8","Ergonómico","Altura entre peldaños cómoda de subir, también para personas mayores","I"),
("R9","Ergonómico","Abrirse y cerrarse con una mano y sin riesgo de pillarse los dedos","I"),
("R10","Ergonómico","Ser ligero para poder llevarlo con una mano","I"),
("R11","Ergonómico","Disponer de un punto de agarre al subir","D"),
("R12","De uso","Ocupar poco espacio una vez plegado","I"),
("R13","De uso","No marcar el suelo ni hacer ruido","I"),
("R14","De uso","No necesitar mantenimiento","D"),
("R15","De uso","Piezas de desgaste (tacos, recubrimientos) sustituibles","D"),
("R16","Estético","Transmitir robustez y seguridad","I"),
("R17","Estético","Aspecto sencillo que no desentone en la vivienda","D"),
("R18","Estético","Disponible en varios acabados de color","D"),
("R19","Medioambiental","Fabricado con materiales reciclables","I"),
("R20","Medioambiental","Larga vida útil","I"),
("R21","Medioambiental","Fácil de desmontar para separar los materiales al final de su vida","D"),
("R22","Medioambiental","Embalaje mínimo y reciclable","D"),
("R23","Económico","Precio de venta competitivo","I")]
tabla(["Código","Tipo","Requisito","I / D"],req,[1.5,3,9.5,1.4],"Tabla 1. Requisitos del producto. Elaboración propia.")

# ---------- 1.6 ----------
H("1.6. PROPIEDADES, ASPECTO Y FORMA",2)
d.add_heading("1.6.1. Propiedades",3)
P("Las propiedades traducen los requisitos anteriores a valores concretos y medibles, que servirán luego para comprobar si el diseño es correcto. Se dividen en pasivas (relacionadas con la estructura y los materiales) y activas (relacionadas con el movimiento y el funcionamiento). Los valores son objetivos de diseño y se revisarán cuando se complete la etapa de información (competencia, norma y EDP).")
prop=[("P1","R1","Carga máxima de uso: 150 kg, sin deformación permanente en el ensayo de resistencia de la norma","Pasiva"),
("P2","R2, R5","Superar el ensayo de estabilidad de la UNE-EN 14183; base de apoyo más ancha que la plataforma","Pasiva"),
("P3","R6, R8","Altura de la plataforma superior: 0,48 m, con dos peldaños separados unos 0,24 m","Pasiva"),
("P4","R7","Profundidad de peldaño ≥ 200 mm y anchura ≥ 350 mm","Pasiva"),
("P5","R10","Peso total ≤ 5 kg","Pasiva"),
("P6","R12","Espesor plegado ≤ 80 mm","Pasiva"),
("P7","R3, R9","Apertura y cierre en ≤ 3 s con una mano, con bloqueo automático en posición de uso","Activa"),
("P8","R9","Holguras en zonas móviles < 8 mm o > 25 mm, para evitar atrapamientos","Activa"),
("P9","R4, R13","Tacos de apoyo de elastómero no marcante y superficie de peldaño antideslizante","Pasiva"),
("P10","R19, R21","≥ 90 % del peso en materiales reciclables; piezas de un solo material, unidas de forma desmontable","Pasiva"),
("P11","R20","Vida útil ≥ 10 años en uso doméstico","Pasiva"),
("P12","R23","Precio de venta objetivo ≤ 45 € (a contrastar en el estudio de mercado)","—")]
tabla(["Código","Requisito","Propiedad (valor objetivo)","Tipo"],prop,[1.5,2,10,1.9],"Tabla 2. Propiedades del producto y requisito del que derivan. Elaboración propia.")
d.add_heading("1.6.2. Aspecto",3)
P("El aspecto buscado es sencillo y actual, con líneas limpias y colores neutros (blanco, gris antracita o acabado madera), para que el taburete no desentone si se queda a la vista en la cocina o en el salón. Al mismo tiempo tiene que transmitir robustez: el usuario debe percibir, antes de subirse, que el producto es firme y que está bien bloqueado. Las indicaciones de uso y la carga máxima deben verse con claridad pero sin recargar el diseño.")
d.add_heading("1.6.3. Forma",3)
P("La forma tiene que estar supeditada a la función. En posición de uso, el taburete debe tener una base de apoyo más amplia que la parte superior, para que el centro de gravedad del usuario quede siempre dentro de la base y se evite el vuelco. Una vez plegado, debe quedar prácticamente plano para poder guardarse en huecos estrechos. Además, no debe tener aristas vivas ni salientes, y debe contar con una zona que permita cogerlo cómodamente con una mano para transportarlo.")

# ---------- Bibliografía ----------
d.add_heading("Bibliografía",2)
bib_txt+=["AENOR. UNE-EN 14183: Taburetes-escalera."]
for t in bib_txt: d.add_paragraph(t,style="List Number")
d.save("Trabajo_Taburete.docx")
print("ok",len(d.inline_shapes),"imagenes",len(d.tables),"tablas")
