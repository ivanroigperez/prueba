import subprocess, re, pymupdf
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
S="/tmp/claude-0/-home-user-prueba/1e15d777-fea3-54a9-ae9b-0a6049279c2c/scratchpad"
F="Trabajo_Taburete.docx"
norm=lambda s: re.sub(r"\s+"," ",s).strip()
def render():
    subprocess.run(["soffice",f"-env:UserInstallation=file://{S}/lo_profile","--headless","--convert-to","pdf","--outdir",S,F],capture_output=True)
    return pymupdf.open(f"{S}/Trabajo_Taburete.pdf")
RI=re.compile(r"^(Ilustración|Plano) \d+\. "); RT=re.compile(r"^Tabla \d+\. ")
def corto(t): return re.split(r" (Elaboración propia|Fuente:|Fuentes:|Valoración cualitativa|Estimación propia)",t)[0].rstrip(". ")
def listas(d):
    H=[(1 if p.style.name=="Heading 1" else 2,p.text) for p in d.paragraphs if p.style.name in("Heading 1","Heading 2") and p.text.strip() and "\t" not in p.text]
    I=[(2,p.text) for p in d.paragraphs if RI.match(p.text) and "\t" not in p.text]
    T=[(2,p.text) for p in d.paragraphs if RT.match(p.text) and "\t" not in p.text]
    return H,I,T
def paginas(doc,items):
    txt=[norm(p.get_text()) for p in doc]
    k0=next(i for i,t in enumerate(txt) if "1.1.1. CONTEXTO" in t); res=[]; k=k0
    for _,h in items:
        key=norm(h)[:40]
        while k<len(txt) and key not in txt[k]: k+=1
        assert k<len(txt),h; res.append(k+1)
    return res
def construir(pagsH,pagsI,pagsT):
    d=Document(F)
    st=d.styles
    for nom in ("Pie de ilustración","Pie de tabla"):
        if nom not in [s.name for s in st]: st.add_style(nom,WD_STYLE_TYPE.PARAGRAPH).base_style=st["Normal"]
    for p in d.paragraphs:
        if "\t" in p.text: continue
        if RI.match(p.text): p.style=st["Pie de ilustración"]
        elif RT.match(p.text): p.style=st["Pie de tabla"]
    H,I,T=listas(d)
    body=d.element.body; ps=d.paragraphs
    ini=[i for i,p in enumerate(ps) if p.text=="ÍNDICE"][0]
    j=ini+1
    while ps[j]._p.find(".//"+qn("w:br"))is None or ps[j].text.strip(): j+=1
    for p in ps[ini+1:j]: body.remove(p._p)
    salto=d.paragraphs[[i for i,p in enumerate(d.paragraphs) if p.text=="ÍNDICE"][0]+1]
    idx=[p for p in d.paragraphs if p.text=="ÍNDICE"][0]
    def fld(p,t):
        r=OxmlElement("w:r"); f=OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"),t); r.append(f); p._p.append(r)
    def bloque(items,pags,instr,titulo=None,pb=False):
        if titulo:
            h=salto.insert_paragraph_before(titulo,style=idx.style); h.paragraph_format.page_break_before=pb
            if not pb: h.paragraph_format.space_before=Pt(24)
        out=[]
        for i,((lvl,t),pg) in enumerate(zip(items,pags)):
            p=salto.insert_paragraph_before(); pf=p.paragraph_format
            pf.left_indent=Cm(0 if lvl==1 else (0.6 if not titulo else 0)); pf.space_before=Pt(8 if lvl==1 else 0); pf.space_after=Pt(2)
            pf.tab_stops.add_tab_stop(Cm(16),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
            if i==0:
                fld(p,"begin"); r=OxmlElement("w:r"); it=OxmlElement("w:instrText"); it.set(qn("xml:space"),"preserve"); it.text=instr; r.append(it); p._p.append(r); fld(p,"separate")
            sz=11 if lvl==1 else 10
            txt=t if not titulo else corto(t)
            r=p.add_run(txt); r.font.size=Pt(sz); r.bold=(lvl==1)
            r=p.add_run("\t"+str(pg)); r.font.size=Pt(sz); r.bold=(lvl==1)
            out.append(p)
        fld(out[-1],"end")
    bloque(H,pagsH,' TOC \\o "1-2" \\h \\z \\u ')
    bloque(I,pagsI,' TOC \\h \\z \\t "Pie de ilustración;1" ',"ÍNDICE DE ILUSTRACIONES",pb=True)
    bloque(T,pagsT,' TOC \\h \\z \\t "Pie de tabla;1" ',"ÍNDICE DE TABLAS",pb=False)
    d.save(F)
d=Document(F); H,I,T=listas(d)
construir([0]*len(H),[0]*len(I),[0]*len(T))
for it in range(3):
    doc=render(); d=Document(F); H,I,T=listas(d)
    pH,pI,pT=paginas(doc,H),paginas(doc,I),paginas(doc,T)
    construir(pH,pI,pT)
    doc=render(); d=Document(F); H2,I2,T2=listas(d)
    ok=(paginas(doc,H2),paginas(doc,I2),paginas(doc,T2))==(pH,pI,pT)
    print("iteración",it,"estable" if ok else "cambia", len(doc),"págs")
    if ok: break
