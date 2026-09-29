import json, copy
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
d=Document("Trabajo_Taburete.docx")
anchor=[p for p in d.paragraphs if p.text=="Bibliografía"][0]
def ins(text="",style=None,bold=False,italic=False,size=None,center=False,before=None):
    b=before or anchor
    p=b.insert_paragraph_before("",style=style) if style else b.insert_paragraph_before("")
    if text:
        r=p.add_run(text); r.bold=bold; r.italic=italic
        if size: r.font.size=Pt(size)
    if center: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    return p
def fig(path,cap,w,before=None):
    p=ins(center=True,before=before); p.add_run().add_picture(path,width=Cm(w)); ins(cap,italic=True,size=9,center=True,before=before)
def shade(c,col):
    tcPr=c._tc.get_or_add_tcPr(); s=OxmlElement("w:shd"); s.set(qn("w:val"),"clear"); s.set(qn("w:color"),"auto"); s.set(qn("w:fill"),col); tcPr.append(s)
def tabla(cab,filas,ws,cap,fs=8):
    t=d.add_table(rows=1,cols=len(cab)); t.style="Table Grid"
    for i,c in enumerate(cab):
        cl=t.rows[0].cells[i]; cl.text=""; r=cl.paragraphs[0].add_run(c); r.bold=True; r.font.size=Pt(fs+0.5); r.font.color.rgb=RGBColor(255,255,255); shade(cl,"C0392B")
    for f in filas:
        cells=t.add_row().cells
        for i,v in enumerate(f):
            cells[i].text=""; r=cells[i].paragraphs[0].add_run(str(v)); r.font.size=Pt(fs)
    tblPr=t._tbl.tblPr; lay=OxmlElement("w:tblLayout"); lay.set(qn("w:type"),"fixed"); tblPr.append(lay)
    for gc,w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")),ws): gc.set(qn("w:w"),str(int(w/2.54*1440)))
    for row in t.rows:
        for c,w in zip(row.cells,ws): c.width=Cm(w)
        for c in row.cells:
            for p in c.paragraphs: p.paragraph_format.keep_with_next=True
    anchor._p.addprevious(t._tbl); ins(cap,italic=True,size=9,center=True)
B=lambda t: ins(t,style="List Bullet")
Q=json.load(open("rediseno/qfd.json")); AV=json.load(open("rediseno/av.json"))

# ===================== 5. REDISEÑO =====================
ins("5. ETAPA DE REDISEÑO",style="Heading 1")
ins("Una vez elegida la solución, se han aplicado las técnicas de rediseño vistas en la asignatura (QFD, AMFEC y análisis del valor) para detectar los puntos débiles del diseño y proponer mejoras antes de pasar a la fabricación de prototipos. El análisis se ha hecho sobre la alternativa 2 (aluminio), que es la que actualmente cumple todas las propiedades del EDP, aunque las conclusiones son aplicables también a la de madera.")
ins("5.1. DESPLIEGUE DE LA FUNCIÓN CALIDAD (QFD)",style="Heading 2")
ins("El QFD traduce lo que pide el cliente (los «qués») en características técnicas medibles (los «cómos»). En la casa de la calidad (Ilustración 19) se han relacionado los nueve requisitos principales del cliente, con su importancia de 1 a 5, con nueve características técnicas del producto. La intensidad de cada relación se ha puntuado con 9 (fuerte), 3 (media) o 1 (débil), y en el tejado se indican las sinergias y los conflictos entre características. A la derecha se compara la percepción del cliente de nuestro diseño con dos productos de la competencia analizados en el apartado 2.2.")
fig("rediseno/figura_qfd.png","Ilustración 19. Casa de la calidad (QFD) del taburete-escalera. Elaboración propia.",16.5)
orden=sorted(zip(Q["comos"],Q["rel"]),key=lambda x:-x[1])
ins(f"Las características técnicas más importantes resultan ser el peso total ({orden[0][1]:.1f} %), la carga máxima ({dict(orden)['Carga máx. (kg)']:.1f} %), la altura de la plataforma y el espesor plegado (alrededor del 12,5 % cada una) y el número de piezas distintas (12,0 %). Esto confirma que el esfuerzo de diseño debe centrarse en conseguir un producto ligero y compacto sin perder resistencia, que es precisamente la contradicción técnica que se resolvió con TRIZ. El tejado muestra el principal conflicto: aumentar la carga máxima tiende a aumentar el peso, mientras que reducir el peso favorece a la vez el espesor plegado, el coste y la huella de carbono.".replace(".",",",0))
ins("En la comparación con la competencia, nuestro diseño iguala al Hailo K20 en seguridad, altura y comodidad y lo supera claramente en peso y precio, mientras que frente al Altipesa ofrece más seguridad y un peldaño mucho más cómodo. Su punto débil es el medio ambiente, por la elevada huella de carbono del aluminio (apartado 4), lo que refuerza la propuesta de usar aluminio reciclado o pasar a la alternativa de madera.")
ins("5.2. ANÁLISIS MODAL DE FALLOS, EFECTOS Y CRITICIDAD (AMFEC)",style="Heading 2")
ins("Con el AMFEC se han identificado los modos de fallo potenciales del producto, valorando para cada uno la probabilidad de que ocurra (O), su gravedad (G) y la probabilidad de que no se detecte antes de llegar al usuario (D), todos de 1 a 10. El producto de los tres da el índice de prioridad de riesgo (IPR = O × G × D), que indica sobre qué fallos hay que actuar primero. Tras proponer una acción correctora se ha vuelto a estimar el IPR.")
am=[("Bloqueo de apertura","Se desbloquea con el usuario encima","Cierre del taburete y caída","Holgura o desgaste del cierre; mal uso",3,10,5,"Pestillo con doble gesto e indicador visual de bloqueo; ensayo al 100 %",2,10,2),
("Tacos antideslizantes","Desgaste o desprendimiento","Deslizamiento sobre suelo liso","Material blando; uso prolongado",4,8,4,"TPE de mayor dureza y taco encajado con tope; repuesto disponible",2,8,3),
("Plataforma de PP","Rotura o fisura por impacto o fatiga","Caída del usuario","Nervios insuficientes; material sin reforzar",2,9,5,"PP con 30 % de fibra de vidrio y nervios; ensayo de resistencia UNE-EN 14183",1,9,3),
("Pasadores y articulaciones","Holgura progresiva","Inestabilidad y ruido","Desgaste del aluminio en los taladros",4,6,6,"Casquillos de nylon en los taladros",2,6,4),
("Bielas y tirantes","Fisura junto al taladro","Pérdida de apoyo de la plataforma o el peldaño","Poca distancia del taladro al borde",2,8,6,"Aumentar la distancia al borde a 12 mm y redondear los extremos",1,8,4),
("Zona de plegado","Atrapamiento de los dedos al cerrar","Lesión en la mano","Holguras entre 8 y 25 mm",4,6,3,"Holguras < 8 mm o > 25 mm y protector en la articulación",2,6,2),
("Largueros y patas","Deformación permanente","Inclinación del taburete","Sobrecarga por encima de 150 kg",2,8,5,"Marcado visible de la carga máxima; ensayo de tipo",2,8,3),
("Pata trasera plegada","Apertura involuntaria al transportarlo","Golpe al usuario","No hay retención en posición cerrada",4,4,4,"Clip de retención en plegado (idea de la patente US 6 454 050)",2,4,3)]
filas=[]
for c,mf,ef,ca,o,g,dd,acc,o2,g2,d2 in am:
    filas.append((c,mf,ef,ca,o,g,dd,o*g*dd,acc,o2*g2*d2))
filas.sort(key=lambda f:-f[7])
tabla(["Componente","Modo de fallo","Efecto","Causa","O","G","D","IPR","Acción correctora","IPR final"],filas,[2.1,2.3,2.0,2.2,0.7,0.7,0.7,0.9,3.4,1.0],"Tabla 16. AMFEC del taburete-escalera. Elaboración propia.",fs=7)
ins(f"El fallo más crítico es el desbloqueo con el usuario encima (IPR = {filas[0][7]}), por su gravedad máxima, seguido de los tacos antideslizantes y las articulaciones. Con las acciones propuestas, sobre todo un pestillo que necesite dos gestos y un indicador visual de bloqueo (idea tomada de la patente US 6 966 404), todos los IPR bajan por debajo de 100 y el más alto pasa a ser de {max(f[9] for f in filas)}.")
ins("5.3. ANÁLISIS DEL VALOR",style="Heading 2")
ins("El análisis del valor compara la importancia que el cliente da a cada función del producto con lo que cuesta cumplirla, para detectar funciones en las que se gasta más de lo que el cliente valora, o al revés. El coste de cada componente se ha estimado a partir del coste industrial objetivo (apartado 2.1) y se ha repartido entre las funciones a las que contribuye; la importancia de cada función se ha obtenido del QFD y del EDP.")
ce=AV["coste_comp"]
tabla(["Componente","Coste estimado (€)"],[(k,f"{v:.2f}".replace(".",",")) for k,v in ce.items()]+[("TOTAL piezas",f"{sum(ce.values()):.2f}".replace(".",","))],[6,4],"Tabla 17. Coste estimado de cada componente (alternativa 2). Estimación propia a partir del coste objetivo.")
fig("rediseno/figura_analisis_valor.png","Ilustración 20. Análisis del valor: valor para el cliente frente a coste de cada función. Elaboración propia.",16)
ins("Del análisis se deducen las siguientes acciones:")
B("Soportar la carga (33 % del coste frente a un 20 % de valor) y la estética (7 % frente a 4 %) están sobredimensionadas. Se propone optimizar la sección de los largueros y las patas, por ejemplo con un perfil de 1,5 mm de pared en las zonas menos cargadas tras un cálculo por elementos finitos, y simplificar el asidero.")
B("El bloqueo (2 % del coste frente a un 12 % de valor) y el antideslizante (1,5 % frente a 8 %) están infravalorados en el diseño actual. Coincide con el AMFEC: conviene invertir en un buen pestillo de seguridad y en tacos y superficies antideslizantes de calidad, porque el cliente los valora mucho y su coste es pequeño.")
B("Transportar y guardar tiene mucho valor y muy poco coste, gracias al plegado plano y al hueco-asa; es uno de los puntos fuertes del producto y conviene destacarlo en la comunicación comercial.")

# ===================== 6. ANEXOS =====================
ins("6. ANEXOS",style="Heading 1")
ins("6.1. PLANOS",style="Heading 2")
ins("Se adjuntan los planos del producto en formato A3 (archivo «Anexo_Planos.pdf»). A continuación se incluye una reproducción reducida de cada uno.")
planos=[("cad/plano1_conjunto_abierto.png","Plano 1. Conjunto abierto con lista de piezas."),("cad/plano2_conjunto_plegado.png","Plano 2. Conjunto plegado."),
        ("cad/plano3_larguero_pata.png","Plano 3. Larguero delantero y pata trasera."),("cad/plano4_plataforma_peldano.png","Plano 4. Plataforma y peldaño."),
        ("cad/plano5_piezas_varias.png","Plano 5. Travesaños, biela, tirante, pasador y taco.")]
for pth,cap in planos:
    p=ins(center=True); p.add_run().add_break(WD_BREAK.PAGE) if False else None
    p.add_run().add_picture(pth,width=Cm(16)); ins(cap,italic=True,size=9,center=True)

# ===================== PORTADA E ÍNDICE =====================
first=d.paragraphs[0]
def pre(text="",size=None,bold=False,center=True,color=None,italic=False):
    p=first.insert_paragraph_before("")
    if text:
        r=p.add_run(text); r.bold=bold; r.italic=italic
        if size: r.font.size=Pt(size)
        if color: r.font.color.rgb=RGBColor(*color)
    if center: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    return p
for _ in range(3): pre()
pre("UNIVERSIDAD DE LA RIOJA",14,True)
pre("Grado en Ingeniería Mecánica · 4.º curso",11)
pre("Ingeniería Simultánea",11,italic=True)
for _ in range(4): pre()
pre("DISEÑO DE UN TABURETE-ESCALERA",24,True,color=(0xC0,0x39,0x2B))
pre("PLEGABLE DE USO DOMÉSTICO",24,True,color=(0xC0,0x39,0x2B))
pre()
pre("Trabajo de la asignatura",12)
pc=pre(); pc.add_run().add_picture("capturas/sw_lateral_abierto.png",width=Cm(7))
for _ in range(3): pre()
al=pre("Alumno: ",11); r=al.add_run("[nombre y apellidos]"); r.font.size=Pt(11); r.font.highlight_color=WD_COLOR_INDEX.YELLOW
pre("Curso 2026-2027",11)
pb=pre(); pb.add_run().add_break(WD_BREAK.PAGE)
h=first.insert_paragraph_before("ÍNDICE"); h.style=d.styles["Heading 1"]
toc=first.insert_paragraph_before("")
fld=OxmlElement("w:fldSimple"); fld.set(qn("w:instr"),'TOC \\o "1-3" \\h \\z \\u')
rr=OxmlElement("w:r"); tt=OxmlElement("w:t"); tt.text="Haz clic derecho aquí y elige «Actualizar campo» para generar el índice."; rr.append(tt); fld.append(rr); toc._p.append(fld)
pb=first.insert_paragraph_before(""); pb.add_run().add_break(WD_BREAK.PAGE)
# que Word actualice los campos al abrir
st=d.settings.element; uf=OxmlElement("w:updateFields"); uf.set(qn("w:val"),"true"); st.append(uf)
# salto de página antes de cada capítulo
for p in d.paragraphs:
    if p.style.name=="Heading 1" and p.text[:2] in ("2.","3.","4.","5.","6.") or (p.style.name=="Heading 1" and p.text.startswith("1. ETAPA")):
        p.paragraph_format.page_break_before=True
    if p.style.name=="Heading 2" and p.text=="Bibliografía": p.paragraph_format.page_break_before=True
d.save("Trabajo_Taburete.docx"); print("ok")
