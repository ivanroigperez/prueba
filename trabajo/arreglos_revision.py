import copy, re
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
d=Document("Trabajo_Taburete.docx")
def P(): return d.paragraphs
def buscar(ini):
    r=[p for p in P() if p.text.startswith(ini) and "\t" not in p.text]; assert len(r)==1,(ini,len(r)); return r[0]
def poner(p,t):
    p.runs[0].text=t
    for r in p.runs[1:]: r.text=""
def clonar(modelo,t,antes=None,despues=None):
    n=copy.deepcopy(modelo._p)
    (antes._p.addprevious(n) if antes is not None else despues._p.addnext(n))
    q=Paragraph(n,modelo._parent); poner(q,t); return q
def borrar(p): p._p.getparent().remove(p._p)
# ---------- tabla con el formato del documento ----------
t15=[t for t in d.tables if t.rows[0].cells[0].text=="Criterio"][0]
TBLPR=copy.deepcopy(t15._tbl.tblPr)
def tabla(filas,anchos,ancla):
    t=d.add_table(rows=len(filas),cols=len(anchos))
    t._tbl.remove(t._tbl.tblPr); t._tbl.insert(0,copy.deepcopy(TBLPR))
    for gc,w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")),anchos): gc.set(qn("w:w"),str(int(w*567)))
    for i,(row,vals) in enumerate(zip(t.rows,filas)):
        for c,w,v in zip(row.cells,anchos,vals):
            tcPr=c._tc.get_or_add_tcPr(); tw=OxmlElement("w:tcW"); tw.set(qn("w:w"),str(int(w*567))); tw.set(qn("w:type"),"dxa"); tcPr.append(tw)
            p=c.paragraphs[0]
            if i<len(filas)-1: p.paragraph_format.keep_with_next=True
            r=p.add_run(v); r.font.size=Pt(8.5)
            if i==0:
                r.bold=True; r.font.color.rgb=__import__("docx").shared.RGBColor(255,255,255)
                sh=OxmlElement("w:shd"); sh.set(qn("w:val"),"clear"); sh.set(qn("w:color"),"auto"); sh.set(qn("w:fill"),"C0392B"); tcPr.append(sh)
            if i==len(filas)-1 and vals[0] in ("Resultado",): r.bold=True
    ancla._p.addprevious(t._tbl); return t
# ---------- 1. renumerar tablas 11-17 -> 12-18 ----------
for n in range(17,10,-1):
    p=buscar(f"Tabla {n}. "); poner(p,p.text.replace(f"Tabla {n}. ",f"Tabla {n+1}. ",1))
cap_modelo=buscar("Tabla 10. ")
# ---------- 2. TRIZ: principio 8 ----------
poner(buscar("Principio 8, contrapeso: no se ve una aplicación directa"),"Principio 8, contrapeso: no se ha encontrado una aplicación directa en este producto.")
poner(buscar("Principio 8, contrapeso: aprovechar el peso"),"Principio 8, contrapeso: no tiene una aplicación directa. En el boceto se pensó en que el propio peso del usuario mantuviera el taburete abierto, pero eso depende de la geometría del mecanismo (posición respecto al punto muerto) y no del principio de contrapeso; se comprueba en el apartado 3.5.")
# ---------- 3. pestillo en 3.4 ----------
poner(buscar("El modelo todavía no incluye el pestillo"),"El modelo todavía no incluye el pestillo de seguridad propuesto con el método TRIZ (contradicción tercera). Como el mecanismo tiene un único grado de libertad, basta con bloquear una de sus articulaciones para fijar el taburete abierto. En el apartado 3.5 se justifica dónde colocarlo (en la articulación del peldaño) y la carga que tiene que soportar; su diseño de detalle queda como acción correctora del AMFEC (apartado 5.2).")
# ---------- 4. nuevo apartado 3.5 ----------
cap4=buscar("4. ETAPA DE EVALUACIÓN"); h2=buscar("3.4. DISEÑO EN SOLID"); normal=buscar("Solo la alternativa de aluminio cumple la propiedad P5")
bul=buscar("Producción: extracción y transformación")
clonar(h2,"3.5. COMPROBACIONES DE RESISTENCIA Y ESTABILIDAD",antes=cap4).paragraph_format.page_break_before=False
clonar(normal,"Antes de evaluar las alternativas se ha comprobado con cálculos manuales sencillos que el diseño de la alternativa 2 cumple las propiedades P1 (150 kg), P2 (estabilidad) y el requisito R3 (no plegarse con el usuario encima). Se toma una carga de 150 kg (1471 N) y se exige un coeficiente de seguridad mínimo de 2. Materiales: aluminio 6063-T5 (límite elástico ≈ 145 MPa, biblioteca de SolidWorks), acero de los pasadores (≈ 235 MPa) y polipropileno con 30 % de fibra de vidrio (resistencia ≈ 75 MPa, valor típico de catálogo).",antes=cap4)
filas=[("Comprobación","Modelo de cálculo","Resultado","C. S.","Conclusión"),
("Larguero delantero","Carga repartida entre los 4 apoyos (368 N por barra); componente transversal en voladizo de 481 mm; tubo 40×20×2 (W = 2223 mm³)","σ ≈ 22 MPa","6,5","Cumple"),
("Pandeo del larguero","Euler, L = 481 mm, I mín = 14 379 mm⁴","Pcr ≈ 42 kN frente a 368 N","> 100","Cumple"),
("Plataforma de PP sin nervios","Viga biapoyada de 350 mm; carga en una huella de 100 mm; sección en U (tapa de 3 mm y paredes de 30 mm)","σ ≈ 70 MPa","1,1","No cumple"),
("Plataforma y peldaño con 3 nervios de 2 mm","Misma viga con 5 almas; + 96 g en total","σ ≈ 37 MPa","2,0","Cumple, por poco"),
("Pasadores Ø 8 (articulación de la plataforma)","Toda la carga en los 2 pasadores; cortadura simple","τ ≈ 15 MPa","9","Cumple"),
("Aplastamiento del PP en los taladros","736 N sobre pared de 3 mm","σ ≈ 31 MPa","2,4","Poco margen: resalte de 6 mm o casquillo"),
("Vuelco frontal y trasero","Centro de gravedad del usuario a 1,40 m; a 254 mm del apoyo delantero y 242 mm del trasero","Empuje horizontal para volcar ≈ 250 N","—","Cumple (estático)"),
("Vuelco lateral","Semiancho de apoyo de las patas traseras: 200 mm","Empuje para volcar ≈ 210 N (14 % del peso)","—","Dirección más desfavorable"),
("Pestillo en la articulación del peldaño","Usuario en el peldaño: momento de cierre 522 N·m; 2 pasadores a 90 mm del pivote","≈ 710 N por pasador; τ ≈ 14 MPa","9","Cumple")]
ancla=clonar(cap_modelo,"",antes=cap4)
tabla(filas,[3.0,5.6,2.9,1.1,3.4],ancla)
poner(ancla,"Tabla 11. Comprobaciones de resistencia y estabilidad de la alternativa 2 (C. S.: coeficiente de seguridad). Elaboración propia.")
clonar(normal,"El resultado más importante es el del mecanismo. Con el modelo cinemático del apartado 3.4 se ha calculado cómo cambia la altura de la carga si la pata trasera gira 1° con los cuatro pies en el suelo:",antes=cap4)
clonar(bul,"Usuario sobre la plataforma: al abrir 1° más, la plataforma baja 3,9 mm, así que su peso tiende a abrir el taburete, no a cerrarlo. La apertura está limitada porque 1,9° más allá de la posición de uso el peldaño y el tirante quedan alineados (punto muerto) y actúan como tope.",antes=cap4)
clonar(bul,"Usuario sobre el peldaño (al subir o bajar): al abrir 1° más, el peldaño sube 6,2 mm, así que su peso tiende a cerrar el taburete, con un momento de unos 522 N·m. Por eso el pestillo es imprescindible para cumplir R3. Se coloca en la articulación del peldaño porque es la que más gira (4,1° por cada grado de la pata), lo que reduce la fuerza sobre él.",antes=cap4)
clonar(normal,"Con los cambios que salen de estas comprobaciones (polipropileno con fibra de vidrio, tres nervios en la plataforma y en el peldaño y casquillos en los taladros), la masa estimada de la alternativa 2 pasa de 3,17 a unos 3,46 kg. Sigue cumpliendo P5, pero con poco margen; si el prototipo pesara más, se puede bajar la pared de los largueros y las patas a 1,5 mm (σ ≈ 29 MPa, C. S. ≈ 5), lo que ahorra unos 0,44 kg. Estos cálculos no sustituyen a los ensayos de la UNE-EN 14183 sobre prototipo, en especial el de deslizamiento lateral.",antes=cap4)
# ---------- 5. decisión 4.5 ----------
poner(buscar("Sin embargo, la huella de carbono no es el único criterio"),"Sin embargo, la huella de carbono no es el único criterio. Antes de comparar, se ha comprobado que cada alternativa cumple las propiedades imprescindibles del EDP, porque un requisito imprescindible no se puede compensar con una buena nota en otros criterios (Tabla 16).")
cap16=buscar("Tabla 16. Matriz de decisión ponderada")
filas=[("Criterio","Acero","Aluminio + PP","Madera"),
("Peso total, P5 ≤ 3,5 kg (imprescindible)","9,20 kg: no cumple","3,17 kg: cumple","3,96 kg: no cumple"),
("Espesor plegado, P6 ≤ 50 mm","40 mm: cumple","40 mm: cumple","40 mm: cumple"),
("Huella de carbono A–C (kg CO₂-eq)","13,92","27,26","2,58"),
("Beneficio del reciclaje, módulo D (kg CO₂-eq)","−6,87","−13,69","−0,43"),
("Resistencia a la humedad (cocina y baño)","Necesita pintura","Buena (anodizado y PP)","Necesita barniz"),
("Resultado","Descartada (P5)","Válida","Descartada (P5)")]
t15._tbl.getparent().remove(t15._tbl)
tabla(filas,[6.0,3.3,3.4,3.3],cap16)
poner(cap16,"Tabla 16. Comprobación de las alternativas frente a las propiedades imprescindibles del EDP. Elaboración propia.")
b=[p for p in P() if p.style.name=="List Bullet" and p.text.startswith(("La alternativa mejor valorada","La alternativa de aluminio es la única","La alternativa de acero es la más barata"))]
assert len(b)==3
poner(b[0],"Solo la alternativa de aluminio y polipropileno cumple todas las propiedades imprescindibles. La de acero (9,20 kg) y la de madera (3,96 kg) superan el peso máximo de 3,5 kg de P5, que sale del requisito imprescindible R10 (poder llevarlo con una mano), por lo que quedan descartadas.")
poner(b[1],"Su punto débil es la huella de carbono, la mayor de las tres (27,3 kg CO₂-eq), debida casi por completo al aluminio primario. Por eso se especifica aluminio con alto contenido reciclado: con el factor del aluminio secundario (2,12 kg CO₂-eq/kg en lugar de 9,08) la huella baja a unos 11,3 kg CO₂-eq, por debajo de la del acero (13,9).")
poner(b[2],"La madera tiene con diferencia la menor huella (2,6 kg CO₂-eq), pero además de superar el peso, su modelo es una simplificación: usa listones macizos con la misma sección que los perfiles metálicos (40 × 20 mm), que habría que aumentar para soportar 150 kg, y un vaciado de la plataforma que no es fabricable en contrachapado. Queda como línea de trabajo futura si se consigue un diseño en madera verificado por debajo de 3,5 kg.")
poner(buscar("Por todo ello, se propone continuar el desarrollo"),"Por todo ello se elige la alternativa 2, con perfiles de aluminio 6063-T5 reciclado y plataforma y peldaño de polipropileno reforzado con fibra de vidrio. Es la que se ha modelado en detalle, la que se representa en los planos y sobre la que se aplican las técnicas de rediseño del apartado 5.")
# ---------- 6. capítulo 5 ----------
poner(buscar("Una vez comparadas las alternativas, se han aplicado"),"Una vez elegida la solución (alternativa 2), se han aplicado las técnicas de rediseño vistas en la asignatura (QFD, AMFEC y análisis del valor) para detectar sus puntos débiles y proponer mejoras antes de pasar a la fabricación de prototipos.")
borrar(buscar("En la alternativa de madera cambian los fallos"))
p=buscar("En la comparación con la competencia")
poner(p,p.text.replace("lo que refuerza la propuesta de usar aluminio reciclado o pasar a la alternativa de madera.","lo que refuerza la decisión de usar aluminio reciclado."))
am=[t for t in d.tables if t.rows[0].cells[0].text=="Componente" and len(t.columns)==10][0]
for row in am.rows:
    c0=row.cells[0].text
    if c0=="Largueros y patas":
        poner(row.cells[8].paragraphs[0],"Sección comprobada por cálculo (apartado 3.5), ensayo de tipo con 150 kg y marcado de la carga máxima")
    if c0=="Bloqueo de apertura":
        poner(row.cells[8].paragraphs[0],"Pestillo de doble gesto en la articulación del peldaño (apartado 3.5) e indicador visual de bloqueo; ensayo al 100 %")
d.save("Trabajo_Taburete.docx"); print("ok")
