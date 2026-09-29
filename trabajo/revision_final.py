from docx import Document
from docx.oxml.ns import qn
import copy
d=Document("Trabajo_Taburete.docx")
P=d.paragraphs
def buscar(ini):
    r=[p for p in d.paragraphs if p.text.startswith(ini)]
    assert len(r)==1,(ini,len(r)); return r[0]
def poner(p,t):
    p.runs[0].text=t
    for r in p.runs[1:]: r.text=""
def despues(p,t):
    n=copy.deepcopy(p._p); p._p.addnext(n)
    from docx.text.paragraph import Paragraph
    q=Paragraph(n,p._parent); poner(q,t); return q
# embalaje
poner(buscar("Dimensiones aproximadas de la caja"),"Dimensiones aproximadas de la caja, a partir del tamaño plegado medido en el modelo (861 × 450 × 40 mm): 880 × 470 × 60 mm.")
poner(buscar("Paletización:"),"Paletización: en un palé europeo (1200 × 800 mm) las cajas se colocan de canto, con su lado largo (880 mm) en la dirección de 1200 mm. Así caben 13 cajas en el lado de 800 mm y 3 alturas de 470 mm, lo que da 39 unidades por palé con una altura total de unos 1,55 m. Colocadas planas solo cabría una caja por capa (unas 23 unidades por palé), por lo que se elige la disposición de canto. El plegado plano del producto es la clave para reducir el volumen transportado y, con ello, el coste y las emisiones del transporte.")
# montaje
poner(buscar("El producto se entrega completamente montado"),"El producto se entrega completamente montado, por lo que el usuario no tiene que realizar ninguna operación salvo desplegarlo. El montaje se hace en fábrica con la siguiente secuencia, actualizada con el diseño definitivo del apartado 3.4:")
pasos=["1. Corte, taladrado y acabado superficial de los largueros delanteros, las patas traseras y los travesaños; inyección de la plataforma y el peldaño.",
"2. Unión del travesaño del asidero a los dos largueros delanteros y del travesaño trasero a las dos patas traseras, formando dos bastidores.",
"3. Colocación de los tacos antideslizantes en los extremos inferiores de los largueros y de las patas.",
"4. Articulación de la plataforma y del peldaño a los largueros delanteros mediante pasadores.",
"5. Unión de las patas traseras a los largueros delanteros en el pivote superior, por encima de la plataforma.",
"6. Montaje de las bielas de la plataforma, los tirantes del peldaño y el pestillo de seguridad.",
"7. Colocación de las etiquetas de advertencia y carga máxima, y control final: apertura y cierre, comprobación del bloqueo y prueba de carga por muestreo."]
for i,t in enumerate(pasos): poner(buscar(f"{i+1}. " if i else "1. Preparación"),t) if i==0 else None
viejos=["2. Colocación de los tacos","3. Unión de los peldaños","4. Unión de la pata trasera","5. Montaje de los tirantes","6. Colocación de las etiquetas","7. Control final"]
for v,t in zip(viejos,pasos[1:]): poner(buscar(v),t)
# pestillo en 3.4
p=buscar("Con el modelo se han comprobado")
despues(p,"El modelo todavía no incluye el pestillo de seguridad propuesto con el método TRIZ (contradicción tercera). Como el mecanismo tiene un único grado de libertad, basta con bloquear una de sus articulaciones para fijar el taburete abierto; la solución concreta (por ejemplo, un pasador con muelle en la articulación de la plataforma, como en el modelo de utilidad ES 1 070 877 U) se deja para el diseño de detalle y aparece como acción correctora en el AMFEC (apartado 5.2).")
# capitulo 5
poner(buscar("Una vez elegida la solución, se han aplicado"),"Una vez comparadas las alternativas, se han aplicado las técnicas de rediseño vistas en la asignatura (QFD, AMFEC y análisis del valor) para detectar los puntos débiles del diseño y proponer mejoras antes de pasar a la fabricación de prototipos. Los análisis se han hecho sobre la alternativa 2 (aluminio y polipropileno), porque es la única cuyo modelo cumple hoy todas las propiedades del EDP y la única con una estimación de costes por componente. La mayoría de las conclusiones afectan al mecanismo (bloqueo, articulaciones, tacos y plegado), que es común a las tres alternativas, así que son igualmente válidas para la alternativa de madera propuesta en el apartado 4.5; al final del AMFEC se indican los cambios propios de la madera.")
p=buscar("El fallo más crítico es el desbloqueo")
despues(p,"En la alternativa de madera cambian los fallos propios del material: la plataforma de contrachapado es menos sensible al impacto que la de polipropileno, pero sí a la humedad (acción: barniz y ensayo en cámara húmeda), y los taladros de las articulaciones necesitan casquillos metálicos para que la madera no se ovalice con el uso.")
p=buscar("Del análisis se deducen las siguientes acciones")
q=copy.deepcopy(p._p); p._p.addprevious(q)
from docx.text.paragraph import Paragraph
poner(Paragraph(q,p._parent),"La suma del coste de las piezas (13,46 €) supera en casi 1 € los 12,50 € previstos para materiales y transformación en el apartado 2.1 (8,50 € + 4,00 €), lo que confirma que hay que ajustar el diseño para cumplir el coste industrial objetivo.")
# planos con el nombre
imgs={1:"plano1_conjunto_abierto",2:"plano2_conjunto_plegado",3:"plano3_larguero_pata",4:"plano4_plataforma_peldano",5:"plano5_piezas_varias"}
n=0
for i,p in enumerate(d.paragraphs):
    for k,f in imgs.items():
        if p.text.startswith(f"Plano {k}."):
            prev=d.paragraphs[i-1]; blip=prev._p.findall('.//'+qn('a:blip'))[0]
            part=d.part.related_parts[blip.get(qn('r:embed'))]; part._blob=open(f"cad/{f}.png","rb").read(); n+=1
print("planos",n)
d.save("Trabajo_Taburete.docx"); print("ok")
