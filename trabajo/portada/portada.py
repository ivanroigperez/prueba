from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import copy
ROJO=RGBColor(0xC0,0x39,0x2B); OSC=RGBColor(0x1F,0x2A,0x36); GRIS=RGBColor(0x6B,0x72,0x80)
d=Document("Trabajo_Taburete.docx")
body=d.element.body
indice=[p for p in d.paragraphs if p.text.strip()=="ÍNDICE"][0]
# borrar la portada anterior
for el in list(body.iterchildren()):
    if el is indice._p: break
    body.remove(el)
indice.paragraph_format.page_break_before=True

def par(texto="",size=11,color=OSC,bold=False,italic=False,align=WD_ALIGN_PARAGRAPH.CENTER,before=0,after=0,spacing=None,font=None):
    p=indice.insert_paragraph_before(); p.alignment=align
    pf=p.paragraph_format; pf.space_before=Pt(before); pf.space_after=Pt(after); pf.line_spacing=1.0
    if texto:
        r=p.add_run(texto); r.font.size=Pt(size); r.font.color.rgb=color; r.bold=bold; r.italic=italic
        if font: r.font.name=font
        if spacing is not None:
            rpr=r._r.get_or_add_rPr(); sp=OxmlElement("w:spacing"); sp.set(qn("w:val"),str(spacing)); rpr.append(sp)
    return p
def sombrear(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); sh=OxmlElement("w:shd"); sh.set(qn("w:val"),"clear"); sh.set(qn("w:color"),"auto"); sh.set(qn("w:fill"),fill); tcPr.append(sh)
def sin_bordes(t):
    tblPr=t._tbl.tblPr; b=OxmlElement("w:tblBorders")
    for k in ("top","left","bottom","right","insideH","insideV"):
        e=OxmlElement("w:"+k); e.set(qn("w:val"),"nil"); b.append(e)
    tblPr.append(b)
def borde_inferior(p,color="C0392B",sz=12):
    pPr=p._p.get_or_add_pPr(); bd=OxmlElement("w:pBdr"); e=OxmlElement("w:bottom")
    e.set(qn("w:val"),"single"); e.set(qn("w:sz"),str(sz)); e.set(qn("w:space"),"4"); e.set(qn("w:color"),color); bd.append(e); pPr.append(bd)
def texto_celda(cell,lineas):
    cell.paragraphs[0].text=""
    for i,(t,size,color,bold) in enumerate(lineas):
        p=cell.paragraphs[0] if i==0 else cell.add_paragraph()
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(0); p.paragraph_format.space_before=Pt(0)
        r=p.add_run(t); r.font.size=Pt(size); r.font.color.rgb=color; r.bold=bold
def tabla(filas,cols,anchos):
    t=d.add_table(rows=filas,cols=cols); t.alignment=WD_TABLE_ALIGNMENT.CENTER; sin_bordes(t)
    t.autofit=False
    for row in t.rows:
        for c,w in zip(row.cells,anchos): c.width=Cm(w)
    indice._p.addprevious(t._tbl); return t

# 1. franja superior
t=tabla(1,1,[16.0]); c=t.cell(0,0); sombrear(c,"1F2A36")
texto_celda(c,[("",4,OSC,False),("UNIVERSIDAD DE LA RIOJA",13,RGBColor(255,255,255),True),
               ("Grado en Ingeniería Mecánica  ·  Ingeniería Simultánea",9.5,RGBColor(0xC9,0xD1,0xDB),False),("",4,OSC,False)])
# 2. titulo
par("TRABAJO DE ASIGNATURA  ·  CURSO 2026-2027",9,ROJO,True,before=30,after=6,spacing=40)
par("Taburete-escalera plegable",26,OSC,True,after=0)
p=par("de uso doméstico",20,GRIS,False,after=10); borde_inferior(p)
par("Diseño, ciclo de vida y rediseño con la metodología de Ingeniería Simultánea",10.5,GRIS,italic=True,before=6,after=14)
# 3. render
p=par(after=4); p.add_run().add_picture("portada/render_portada.png",height=Cm(9.3))
par("Plegado  ·  Abierto",8.5,GRIS,italic=True,after=14)
# 4. cifras clave
t=tabla(1,4,[4.0,4.0,4.0,4.0])
for c,(v,l) in zip(t.rows[0].cells,(("0,48 m","altura de plataforma"),("150 kg","carga máxima"),("40 mm","espesor plegado"),("3,17 kg","peso (alt. aluminio)"))):
    sombrear(c,"F3F4F6"); texto_celda(c,[("",3,OSC,False),(v,17,ROJO,True),(l,8.5,GRIS,False),("",3,OSC,False)])
# 5. autor
par("",before=30)
p=par("Alumno",8.5,GRIS,False,after=0,spacing=30)
par("Iván Roig Pérez",15,OSC,True,after=2)
par("Logroño, 2026",9.5,GRIS,after=0)

# pie de pagina con numero (sin numero en la portada)
sec=d.sections[0]; sec.different_first_page_header_footer=True
fp=sec.footer.paragraphs[0]; fp.text=""; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
def campo(p,instr):
    r=p.add_run(); f1=OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"),"begin"); r._r.append(f1)
    r=p.add_run(); it=OxmlElement("w:instrText"); it.set(qn("xml:space"),"preserve"); it.text=instr; r._r.append(it)
    r=p.add_run(); f2=OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"),"separate"); r._r.append(f2)
    r=p.add_run("1"); r.font.size=Pt(9); r.font.color.rgb=GRIS
    r=p.add_run(); f3=OxmlElement("w:fldChar"); f3.set(qn("w:fldCharType"),"end"); r._r.append(f3)
r=fp.add_run("Taburete-escalera plegable  ·  Iván Roig Pérez  ·  "); r.font.size=Pt(9); r.font.color.rgb=GRIS
campo(fp," PAGE ")
d.save("Trabajo_Taburete.docx"); print("ok")
