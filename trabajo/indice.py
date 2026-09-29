import subprocess, re, pymupdf, copy
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
S="/tmp/claude-0/-home-user-prueba/1e15d777-fea3-54a9-ae9b-0a6049279c2c/scratchpad"
def render():
    subprocess.run(["soffice",f"-env:UserInstallation=file://{S}/lo_profile","--headless","--convert-to","pdf","--outdir",S,"Trabajo_Taburete.docx"],capture_output=True)
    return pymupdf.open(f"{S}/Trabajo_Taburete.pdf")
norm=lambda s: re.sub(r"\s+"," ",s).strip()
def paginas(doc,heads):
    txt=[norm(p.get_text()) for p in doc]; res=[]; k=2
    for lvl,h in heads:
        key=norm(h)[:35]
        while k<len(txt) and key not in txt[k]: k+=1
        assert k<len(txt),h; res.append(k+1)
    return res
d=Document("Trabajo_Taburete.docx")
heads=[(1 if p.style.name=="Heading 1" else 2,p.text) for p in d.paragraphs if p.style.name in("Heading 1","Heading 2") and p.text.strip()]
pags=paginas(render(),heads)
# construir el campo TOC con su resultado ya calculado
viejo=[p for p in d.paragraphs if p._p.find(qn('w:fldSimple')) is not None and 'TOC' in p._p.find(qn('w:fldSimple')).get(qn('w:instr'))]
ancla=viejo[0]
def fld(p,tipo):
    r=OxmlElement('w:r'); f=OxmlElement('w:fldChar'); f.set(qn('w:fldCharType'),tipo); r.append(f); p._p.append(r)
nuevos=[]
for i,((lvl,h),pg) in enumerate(zip(heads,pags)):
    p=ancla.insert_paragraph_before(); pf=p.paragraph_format
    pf.left_indent=Cm(0 if lvl==1 else 0.6); pf.space_before=Pt(8 if lvl==1 else 0); pf.space_after=Pt(2)
    pf.tab_stops.add_tab_stop(Cm(16),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
    if i==0:
        fld(p,'begin'); r=OxmlElement('w:r'); it=OxmlElement('w:instrText'); it.set(qn('xml:space'),'preserve'); it.text=' TOC \\o "1-2" \\h \\z \\u '; r.append(it); p._p.append(r); fld(p,'separate')
    r=p.add_run(h); r.font.size=Pt(11 if lvl==1 else 10); r.bold=(lvl==1)
    r=p.add_run("\t"+str(pg)); r.font.size=Pt(11 if lvl==1 else 10); r.bold=(lvl==1)
    nuevos.append(p)
fld(nuevos[-1],'end')
ancla._p.getparent().remove(ancla._p)
d.save("Trabajo_Taburete.docx")
# comprobar que las paginas no han cambiado al meter el indice
d2=Document("Trabajo_Taburete.docx")
heads2=[(1,p.text) for p in d2.paragraphs if p.style.name in("Heading 1","Heading 2") and p.text.strip()]
pags2=paginas(render(),heads2)
print(len(heads),"entradas; cambios:",[(h,a,b) for (_,h),a,b in zip(heads,pags,pags2) if a!=b])
